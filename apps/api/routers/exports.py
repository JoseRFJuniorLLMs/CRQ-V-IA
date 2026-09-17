import csv
import io
from io import BytesIO
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, Query, Request, Response
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import or_, desc
import openpyxl
from apps.api.core.database import get_db
from apps.api.core.deps import get_current_user, record_audit
from apps.api.models.user import User
from apps.api.models.company import Establishment, Company
from apps.api.models.saved_list import SavedList

router = APIRouter(prefix="/exports", tags=["Exportações"])

def build_prospect_query(db: Session, q: str = None, city: str = None, cnae: str = None, tier: str = None, status: str = None):
    query = db.query(Establishment).join(Company)
    if q:
        clean_q = q.strip()
        query = query.filter(or_(Company.legal_name.ilike(f"%{clean_q}%"), Establishment.trade_name.ilike(f"%{clean_q}%"), Establishment.cnpj.ilike(f"%{clean_q}%")))
    if city:
        query = query.filter(Establishment.city.ilike(f"%{city}%"))
    if cnae:
        query = query.filter(Establishment.primary_cnae.ilike(f"%{cnae}%"))
    if tier:
        query = query.filter(Establishment.cfq_tier == tier.upper())
    if status:
        query = query.filter(Establishment.registration_status == status.upper())
    return query.order_by(desc(Establishment.chemical_score))

@router.get("/prospects.csv")
def export_prospects_csv(
    request: Request,
    q: Optional[str] = None,
    city: Optional[str] = None,
    cnae: Optional[str] = None,
    tier: Optional[str] = None,
    status: Optional[str] = None,
    limit: int = Query(500, le=2000),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    query = build_prospect_query(db, q, city, cnae, tier, status)
    records = query.limit(limit).all()

    output = io.StringIO()
    writer = csv.writer(output, delimiter=";", quoting=csv.QUOTE_MINIMAL)
    writer.writerow(["CNPJ", "Razao_Social", "Nome_Fantasia", "Municipio", "UF", "CNAE_Principal", "Situacao_RFB", "Porte", "Score_Quimica", "Prioridade_CFQ", "Justificativa_Regulatoria", "Status_CRQV"])

    for est in records:
        writer.writerow([
            est.cnpj,
            est.company.legal_name,
            est.trade_name or "",
            est.city,
            est.state,
            est.primary_cnae,
            est.registration_status,
            est.company.company_size,
            f"{est.chemical_score:.1f}",
            est.cfq_tier,
            est.cfq_rationale or "",
            est.crq_status
        ])

    output.seek(0)
    client_ip = request.client.host if request.client else "127.0.0.1"
    record_audit(
        db=db,
        action="EXPORT_PROSPECTS_CSV",
        user=current_user,
        target_type="EXPORT",
        details={"records_count": len(records), "city": city, "tier": tier},
        ip_address=client_ip
    )

    return Response(
        content=output.getvalue().encode("utf-8-sig"),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=prospeccao_crqv_rs.csv"}
    )

@router.get("/prospects.xlsx")
def export_prospects_xlsx(
    request: Request,
    q: Optional[str] = None,
    city: Optional[str] = None,
    cnae: Optional[str] = None,
    tier: Optional[str] = None,
    status: Optional[str] = None,
    limit: int = Query(500, le=2000),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    query = build_prospect_query(db, q, city, cnae, tier, status)
    records = query.limit(limit).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Prospecção CRQ-V"

    headers = ["CNPJ", "Razão Social", "Nome Fantasia", "Município", "UF", "CNAE Principal", "Situação RFB", "Porte", "Score Química", "Prioridade CFQ", "Enquadramento CFQ 339/2025", "Status no CRQ-V"]
    ws.append(headers)

    for est in records:
        ws.append([
            est.cnpj,
            est.company.legal_name,
            est.trade_name or "",
            est.city,
            est.state,
            est.primary_cnae,
            est.registration_status,
            est.company.company_size,
            est.chemical_score,
            est.cfq_tier,
            est.cfq_rationale or "",
            est.crq_status
        ])

    # Ajusta largura de colunas
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 40)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    client_ip = request.client.host if request.client else "127.0.0.1"
    record_audit(
        db=db,
        action="EXPORT_PROSPECTS_XLSX",
        user=current_user,
        target_type="EXPORT",
        details={"records_count": len(records)},
        ip_address=client_ip
    )

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=prospeccao_crqv_rs.xlsx"}
    )

@router.get("/lists/{list_id}.xlsx")
def export_list_xlsx(
    list_id: int,
    request: Request,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    saved_list = db.query(SavedList).filter(SavedList.id == list_id).first()
    if not saved_list:
        raise HTTPException(status_code=404, detail="Lista não encontrada")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Roteiro {saved_list.name[:20]}"

    headers = ["CNPJ", "Razão Social", "Nome Fantasia", "Município", "Endereço", "Telefone", "CNAE", "Score", "Prioridade", "Status Fiscal", "Notas do Fiscal"]
    ws.append(headers)

    for item in saved_list.items:
        est = item.establishment
        if est:
            address_str = f"{est.street or ''}, {est.number or ''} {est.complement or ''} - {est.district or ''}"
            ws.append([
                est.cnpj,
                est.company.legal_name,
                est.trade_name or "",
                est.city,
                address_str,
                est.phone or "",
                est.primary_cnae,
                est.chemical_score,
                item.priority,
                item.fiscal_status,
                item.notes or ""
            ])

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    client_ip = request.client.host if request.client else "127.0.0.1"
    record_audit(
        db=db,
        action="EXPORT_SAVED_LIST_XLSX",
        user=current_user,
        target_type="SAVED_LIST",
        target_id=str(list_id),
        details={"list_name": saved_list.name, "items_count": len(saved_list.items)},
        ip_address=client_ip
    )

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=roteiro_fiscalizacao_{list_id}.xlsx"}
    )
